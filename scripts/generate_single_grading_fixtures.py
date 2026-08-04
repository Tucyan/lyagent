from __future__ import annotations

import json
from pathlib import Path
from shutil import copyfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "tests" / "fixtures" / "single-grading" / "ai-life-report"
SOURCE_RUBRIC = Path(r"D:\Users\almerb\Downloads\export.json")
TITLE = "生成式 AI 生活助手的设计与实现报告"
STUDENT = "张晓明（20260001）"
REPORT_MD = """# 生成式 AI 生活助手的设计与实现报告

**学生：张晓明（20260001）**

## 一、问题背景、待解决问题与意义

生成式 AI 已进入学习、健康信息整理和家庭事务管理等日常场景，但通用聊天工具经常给出过长、缺少来源边界或无法直接执行的建议。本项目要解决的问题是：如何把分散的自然语言需求转换为结构化、可复核的生活行动清单，同时明确提示隐私风险与人工判断边界。

该问题具有现实意义。学生可以更快整理课程任务，家庭成员可以统一购物和日程信息，使用者也能在健康咨询前准备问题清单。系统不替代教师、医生或其他专业人员，而是提供信息整理和决策准备支持。

## 二、成果的详细描述

我完成了一个名为“生活助手”的网页原型，包含需求输入、场景选择、结果预览和历史记录四个区域。用户可选择学习计划、就医准备或家庭采购场景，输入目标、截止时间与限制条件，系统会生成分步骤清单，并在涉及健康和隐私时显示醒目提醒。

学习计划场景会按优先级输出任务、预计时长和复习检查点；就医准备场景只整理症状记录与待询问事项，不提供诊断；家庭采购场景可合并重复物品并按类别排序。历史记录仅保存在本地浏览器，用户可以逐条删除。界面使用清晰标题、状态提示和确认按钮，错误输入会得到可操作的修正说明。

| 场景 | 主要输入 | 主要结果 | 风险控制 |
|---|---|---|---|
| 学习计划 | 任务、期限、每日可用时间 | 优先级清单与检查点 | 提醒教师要求优先 |
| 就医准备 | 症状记录、持续时间 | 就诊问题清单 | 明示不作诊断 |
| 家庭采购 | 成员需求、预算 | 分类采购清单 | 敏感信息不上传 |

## 三、开发实现过程

### 3.1 需求拆分与结构设计

我先把全部功能拆为输入校验、提示词组装、模型请求、结果解析、风险提醒和本地历史六个步骤。数据结构使用场景、目标、约束和条目数组，避免模型输出直接控制页面。每个场景共用同一条处理管线，只在提示词模板与结果字段上做小范围差异。

### 3.2 输入提示词

学习计划功能使用的核心提示词如下：

> 你是学习任务整理助手。根据用户提供的任务、截止日期和每日可用时间，返回严格 JSON。每个条目必须包含任务名、优先级、预计分钟数和检查点。不要改变教师要求；信息不足时列出待确认项。

就医准备功能使用的核心提示词如下：

> 你只负责整理就医前信息，不诊断、不推荐药物。根据症状描述生成时间线、待补充记录和可向医生询问的问题。遇到紧急危险信号时提醒用户立即联系当地急救服务。

家庭采购功能使用的核心提示词如下：

> 将家庭成员提交的采购需求合并为分类清单，保留数量与预算限制，标记冲突项。不要输出输入中不存在的敏感个人信息。

### 3.3 系统输出与联调记录

本次提交未保留模型的实际系统输出，只记录了“返回 JSON 并成功渲染”的结论。因此，读者无法核对条目字段、风险提示文字以及异常输入时的真实返回内容。这是报告证据链的明确缺口，但输入提示词、全部功能步骤与页面处理逻辑均已记录。

### 3.4 验证方式

我分别用正常输入、缺少截止日期、超长文本和包含健康风险描述的输入进行检查。前端先执行长度与必填校验，服务端再验证 JSON 字段；解析失败时显示重试建议，不把原始模型响应直接注入页面。六个实现步骤均通过人工走查，三个场景也都覆盖了输入、请求、解析、展示和历史删除流程。

## 四、问题、讨论与改进方案

目前最大的不足是系统输出证据没有随报告保存，使提示词与最终界面之间缺少可审计的中间记录。产生这一问题的原因是开发时只关注界面验收，没有把匿名化样例输出纳入交付清单。下一版将为每个场景保存一组去标识化的请求与响应样例，并记录模型版本、时间、解析结果和失败分支；同时增加自动测试，确保风险提示字段不可缺失。

第二个问题是历史记录仅依赖浏览器本地存储，跨设备无法同步。后续可以在获得明确同意后提供加密同步，并设置自动过期时间；健康相关内容默认不同步。第三个问题是预计时长由模型生成，可能显得精确但依据不足。改进方式是把时长改成区间，要求用户确认，并在多次完成任务后用个人历史数据校准。

## 五、总结

本项目完成了三个生活场景的生成式 AI 助手原型，并通过受控数据结构、输入校验和风险提醒限制模型能力边界。成果能够把自然语言需求转成可执行清单，也明确暴露了系统输出证据缺失这一问题。后续工作的优先级是补齐可审计样例输出，其次是完善隐私友好的同步机制和时长估计方法。
"""

def set_font(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None: tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(dxa)); tc_w.set(qn("w:type"), "dxa")
    if tc_w.getparent() is None: tc_pr.append(tc_w)

def build_docx(path: Path):
    doc = Document(); section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.492)
    normal = doc.styles["Normal"]; normal.font.name = "Microsoft YaHei"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei"); normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8); normal.paragraph_format.line_spacing = 1.333
    for style_name, size, color, before, after in [("Heading 1",16,"2E74B5",18,10),("Heading 2",13,"2E74B5",12,6),("Heading 3",12,"1F4D78",8,4)]:
        style=doc.styles[style_name]; style.font.name="Microsoft YaHei"; style._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); style.font.size=Pt(size); style.font.color.rgb=RGBColor.from_string(color); style.paragraph_format.space_before=Pt(before); style.paragraph_format.space_after=Pt(after)
    header = section.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.LEFT; set_font(header.add_run("课程作业报告 · 生成式 AI 与生活"), 9, color="68717E")
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT; set_font(footer.add_run("张晓明 · 20260001"), 9, color="68717E")
    for _ in range(5): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(10); set_font(p.add_run(TITLE),30,True,"203748")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(70); set_font(p.add_run("课程项目报告"),15,False,"2B5163")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run(STUDENT),12,True,"203748")
    doc.add_page_break()
    lines=REPORT_MD.splitlines()[4:]; i=0
    while i < len(lines):
        line=lines[i]
        if line.startswith("## "): doc.add_heading(line[3:], level=1)
        elif line.startswith("### "): doc.add_heading(line[4:], level=2)
        elif line.startswith("> "):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.right_indent=Inches(.2); set_font(p.add_run(line[2:]),10.5,False,"444444")
        elif line.startswith("| 场景"):
            rows=[]
            while i < len(lines) and lines[i].startswith("|"):
                if "---" not in lines[i]: rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            table=doc.add_table(rows=len(rows), cols=4); table.alignment=WD_TABLE_ALIGNMENT.LEFT; table.autofit=False
            widths=[1600,2500,2860,2400]
            for r,row in enumerate(rows):
                for c,value in enumerate(row):
                    cell=table.cell(r,c); set_cell_width(cell,widths[c]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text=value
                    for run in cell.paragraphs[0].runs: set_font(run,9.5,r==0)
                    if r==0: cell._tc.get_or_add_tcPr().append(OxmlElement("w:shd")); cell._tc.tcPr[-1].set(qn("w:fill"),"F4F6F9")
            tbl_pr=table._tbl.tblPr; tbl_w=tbl_pr.find(qn("w:tblW")); tbl_w.set(qn("w:w"),"9360"); tbl_w.set(qn("w:type"),"dxa")
            i -= 1
        elif line and not line.startswith("**学生"):
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; set_font(p.add_run(line),11)
        i += 1
    doc.save(path)

def build_pdf(path: Path):
    font_path=Path(r"C:\Windows\Fonts\msyh.ttc")
    pdfmetrics.registerFont(TTFont("YaHei", str(font_path), subfontIndex=0))
    styles=getSampleStyleSheet(); body=ParagraphStyle("cn",parent=styles["BodyText"],fontName="YaHei",fontSize=10.5,leading=15,spaceAfter=8,alignment=TA_JUSTIFY); h1=ParagraphStyle("h1cn",parent=body,fontSize=16,leading=20,textColor=colors.HexColor("#2E74B5"),spaceBefore=14,spaceAfter=8); h2=ParagraphStyle("h2cn",parent=body,fontSize=13,leading=17,textColor=colors.HexColor("#1F4D78"),spaceBefore=10,spaceAfter=6)
    story=[Spacer(1,1.5*inch),Paragraph(TITLE,ParagraphStyle("title",parent=body,fontSize=28,leading=35,alignment=TA_CENTER,textColor=colors.HexColor("#203748"))),Spacer(1,.25*inch),Paragraph("课程项目报告",ParagraphStyle("sub",parent=body,fontSize=15,alignment=TA_CENTER,textColor=colors.HexColor("#2B5163"))),Spacer(1,1.25*inch),Paragraph(STUDENT,ParagraphStyle("student",parent=body,fontSize=12,alignment=TA_CENTER)),PageBreak()]
    lines=REPORT_MD.splitlines()[4:]; i=0
    while i<len(lines):
        line=lines[i]
        if line.startswith("## "): story.append(Paragraph(line[3:],h1))
        elif line.startswith("### "): story.append(Paragraph(line[4:],h2))
        elif line.startswith("> "): story.append(Table([[Paragraph(line[2:],body)]],colWidths=[6.25*inch],style=TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#F4F6F9")),("BOX",(0,0),(-1,-1),.5,colors.HexColor("#D9DEE6")),("LEFTPADDING",(0,0),(-1,-1),10),("RIGHTPADDING",(0,0),(-1,-1),10),("TOPPADDING",(0,0),(-1,-1),8),("BOTTOMPADDING",(0,0),(-1,-1),8)])))
        elif line.startswith("| 场景"):
            rows=[]
            while i<len(lines) and lines[i].startswith("|"):
                if "---" not in lines[i]: rows.append([Paragraph(c.strip(),ParagraphStyle("cell",parent=body,fontSize=8.5,leading=11)) for c in lines[i].strip("|").split("|")])
                i+=1
            story.append(Table(rows,colWidths=[1.05*inch,1.65*inch,1.9*inch,1.65*inch],repeatRows=1,style=TableStyle([("GRID",(0,0),(-1,-1),.5,colors.HexColor("#CBD2DB")),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#F4F6F9")),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5)]))); i-=1
        elif line and not line.startswith("**学生"): story.append(Paragraph(line,body))
        i+=1
    def furniture(canvas, doc):
        canvas.saveState(); canvas.setFont("YaHei",8); canvas.setFillColor(colors.HexColor("#68717E")); canvas.drawString(inch,10.55*inch,"课程作业报告 · 生成式 AI 与生活"); canvas.drawRightString(7.5*inch,.45*inch,f"第 {doc.page} 页"); canvas.restoreState()
    SimpleDocTemplate(str(path),pagesize=letter,rightMargin=inch,leftMargin=inch,topMargin=.75*inch,bottomMargin=.7*inch,title=TITLE,author=STUDENT).build(story,onFirstPage=furniture,onLaterPages=furniture)

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    copyfile(SOURCE_RUBRIC,OUT/"rubric-v1.json")
    (OUT/"20260001_张晓明_生成式AI生活助手报告.md").write_text(REPORT_MD,encoding="utf-8")
    build_docx(OUT/"20260001_张晓明_生成式AI生活助手报告.docx")
    build_pdf(OUT/"20260001_张晓明_生成式AI生活助手报告.pdf")
    parsed=json.loads((OUT/"rubric-v1.json").read_text(encoding="utf-8"))
    assert parsed["hash"]=="949736269815dc30601a22a1fb1af84aee4752b11166685eb278e40becec4ae7"

if __name__ == "__main__": main()
